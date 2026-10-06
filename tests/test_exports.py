import sys
import io
import csv
sys.path.append("backend")

from starlette.testclient import TestClient
from app.main import app
import openpyxl

def test_session_report_csv_and_excel_exports():
    client = TestClient(app)

    # 1. Resolve active demo session
    res = client.get("/api/sessions/active")
    assert res.status_code == 200
    session_id = res.json()["id"]
    print(f"[TEST EXPORT] Active session: {session_id}")

    # 2. Test CSV Export Endpoint
    csv_res = client.get(f"/api/sessions/{session_id}/export/csv")
    assert csv_res.status_code == 200, f"CSV export failed: {csv_res.text}"
    assert "text/csv" in csv_res.headers["content-type"]
    assert "Content-Disposition" in csv_res.headers
    assert ".csv" in csv_res.headers["Content-Disposition"]

    # Verify CSV content and BOM
    raw_csv = csv_res.content
    assert raw_csv.startswith(b"\xef\xbb\xbf"), "Missing UTF-8 BOM for Excel compatibility"

    text_csv = raw_csv.decode("utf-8")
    lines = list(csv.reader(io.StringIO(text_csv)))
    assert len(lines) > 5, f"CSV output too short: {lines}"
    
    # Check for metadata and header rows
    header_found = any("Student Name" in row for row in lines)
    assert header_found, "CSV missing 'Student Name' table header"
    print(f"[TEST EXPORT] CSV export verified successfully! Total rows: {len(lines)}")

    # 3. Test Excel (.xlsx) Export Endpoint
    excel_res = client.get(f"/api/sessions/{session_id}/export/excel")
    assert excel_res.status_code == 200, f"Excel export failed: {excel_res.text}"
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in excel_res.headers["content-type"]
    assert ".xlsx" in excel_res.headers["Content-Disposition"]
    assert len(excel_res.content) > 1000

    # Verify openpyxl can load the workbook
    wb = openpyxl.load_workbook(io.BytesIO(excel_res.content))
    assert "Attendance & Engagement" in wb.sheetnames
    ws = wb["Attendance & Engagement"]
    assert ws["A1"].value is not None
    assert "STUDENT ENGAGEMENT" in str(ws["A1"].value)
    
    # Check that headers row exists
    header_row_values = [ws.cell(row=8, column=col).value for col in range(1, 12)]
    assert "Student Name" in header_row_values
    assert "Engagement Score" in header_row_values

    print(f"[TEST EXPORT] Excel (.xlsx) verified successfully! Sheet: {ws.title}, max_row: {ws.max_row}")
    print("\n>>> ALL EXPORT TESTS PASSED SUCCESSFULLY! 100% VERIFIED!")

if __name__ == "__main__":
    test_session_report_csv_and_excel_exports()
